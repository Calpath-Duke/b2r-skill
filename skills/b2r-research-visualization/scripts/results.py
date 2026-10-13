"""Visualization of supplied numeric tables; no hypothesis tests or bioinformatics."""
from __future__ import annotations

import csv
import io
import math
import re
import shutil
import statistics
from collections import Counter
from pathlib import Path

from common import (require, fields, finite, text_value, color, CliError, checked_input_file,
                    matplotlib_setup, write_json, export, text_checks, finish, pending, digest)
from color_palettes import OKABE_ITO_ON_WHITE

TIME={'raw_timecourse_sd','raw_timecourse_sem'}
MODES={'raw_points','raw_mean_sd','raw_mean_sem','raw_box','raw_paired','summary_point','summary_bar','line'}|TIME
UNIT_MODES={'raw_mean_sd','raw_mean_sem','raw_box','raw_paired'}|TIME
WIDE_MODES={'raw_points','raw_mean_sd','raw_mean_sem','raw_box','raw_paired'}
MAX_GROUPS=8
PALETTE=list(OKABE_ITO_ON_WHITE)+['#56B4E9','#E69F00','#7F7F7F']
MARKS=['o','s','^','D','v','P','X','<']
HATCH=['','//','..','xx','--','\\\\','++','oo']
LINES=['-','--','-.',':']*2
INK='#2B3640'
DEFAULT_MISSING=['','NA','N/A']


def load_table(path, sheet=None):
    p=checked_input_file(path,max_bytes=20_000_000)
    if p.suffix.lower() in ['.csv','.tsv']:
        require(sheet is None,'CSV/TSV 没有工作表；去掉 sheet')
        rows=list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig')),
                             delimiter='\t' if p.suffix.lower()=='.tsv' else ','))
    elif p.suffix.lower()=='.xlsx':
        import zipfile
        with zipfile.ZipFile(p) as z:
            require(sum(i.file_size for i in z.infolist())<=80_000_000,'XLSX 解压内容过大；请导出 CSV')
        try:
            import openpyxl
        except ImportError:
            raise CliError('XLSX 需要可选依赖 openpyxl；也可另存为 CSV') from None
        wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
        try:
            names=wb.sheetnames
            if sheet is None:
                require(len(names)==1,'XLSX 有多个工作表：'+'、'.join(names[:12])+'；请用 sheet 指定其一')
                ws=wb[names[0]]
            else:
                require(sheet in names,'工作表不存在；可选：'+'、'.join(names[:12]))
                ws=wb[sheet]
            rows=[list(r) for r in ws.iter_rows(values_only=True)]
        finally: wb.close()
    else: raise CliError('输入只支持 CSV、TSV、XLSX')
    while rows and all(c is None or str(c).strip()=='' for c in rows[-1]): rows.pop()
    require(2<=len(rows)<=20001,'表格需要标题和 1—20000 行数据')
    header=[str(c).strip() if c is not None else '' for c in rows[0]]
    while header and not header[-1] and all(len(r)<len(header) or r[len(header)-1] in (None,'') for r in rows[1:]):
        header.pop()
    require(all(header) and len(header)==len(set(header)),'列名为空或重复')
    require(len(header)<=100,'列数最多 100')
    result=[]
    for i,row in enumerate(rows[1:],2):
        row=list(row)
        if not row: row=['']*len(header)
        if len(row)>len(header) and all(c in (None,'') for c in row[len(header):]): row=row[:len(header)]
        require(len(row)==len(header),f'第 {i} 行列数不一致')
        result.append({k:'' if v is None else str(v) for k,v in zip(header,row)})
    return header,result


def number(raw, name, missing_values, allow_missing=True):
    if raw.strip() in missing_values:
        require(allow_missing,f'{name} 不允许缺失')
        return None
    try: v=float(raw)
    except (ValueError,TypeError): raise CliError(f'{name} 不是数字') from None
    require(math.isfinite(v),f'{name} 含非有限数；请明确缺失值编码')
    return v


def validate(spec, header):
    fields(spec,['kind','title','purpose','language','mode','columns','unit','xlabel','ylabel',
                 'group_order','missing_values','transform','error_definition','replication_unit',
                 'independence_confirmed','width_mm','height_mm','colors','notes',
                 'layout','design','technical_replicates','sheet'], '结果图')
    require(spec.get('kind')=='results','kind 必须是 results')
    mode=spec.get('mode')
    require(mode in MODES, '不支持 mode；请参考 input-contracts.md')
    for k in ['title','purpose','unit']: text_value(spec.get(k),k,180)
    require(spec.get('language') in ['zh','en','bilingual'],'language 必须是 zh/en/bilingual')
    layout=spec.get('layout','long')
    require(layout in ['long','wide'],'layout 只能是 long/wide')
    cols=spec.get('columns',{})
    fields(cols,['group','value','sample_id','x','error','lower','upper','n','groups'],'columns')
    if layout=='wide':
        require(mode in WIDE_MODES,'wide 只用于原始观测；已有汇总和时间序列请用长表（每行一个观测）')
        groups=cols.get('groups')
        require(isinstance(groups,list) and groups and len(groups)==len(set(groups))
                and all(isinstance(c,str) and c in header for c in groups),'wide 需要 columns.groups：作为分组的列名列表')
        require(set(cols)<={'groups','sample_id'},'wide 只接受 columns.groups 和可选 sample_id')
        require(cols.get('sample_id') not in groups,'sample_id 不能同时是分组列')
        singles={k:v for k,v in cols.items() if k!='groups'}
    else:
        require('groups' not in cols,'columns.groups 只用于 layout: wide')
        require('group' in cols and 'value' in cols,'columns 需要 group 和 value')
        singles=cols
    require(all(isinstance(v,str) and v in header for v in singles.values()),'指定列不存在')
    require(len(set(singles.values()))==len(singles),'多个变量不能映射同一列')
    raw=mode.startswith('raw_')
    require(not(raw and set(cols)&{'error','lower','upper','n'}),'raw 模式不接收已有误差、n')
    if mode in TIME:
        require('x' in cols and 'sample_id' in cols,'重复测量时间曲线需要数值 x 列和受试单位 sample_id 列')
    else:
        require(not(raw and 'x' in cols),'raw 模式不接收 x；逐个体时间曲线用 raw_timecourse_sd/sem')
    require(not('error' in cols and set(cols)&{'lower','upper'}),'对称与上下界误差不能混用')
    require(('lower' in cols)==('upper' in cols),'上下界必须同时提供')
    if mode=='line': require('x' in cols,'line 模式需要数值 x 列')
    if mode in {'summary_point','summary_bar'}: require('x' not in cols,'分组汇总模式不接受 x')
    require(spec.get('transform','none') in ['none','log2'],'transform 仅支持 none/log2')
    tech=spec.get('technical_replicates','none')
    require(tech in ['none','mean'],'technical_replicates 只能是 none/mean')
    if tech=='mean': require(raw and 'sample_id' in cols,'technical_replicates: mean 需要原始模式和 sample_id（生物学重复单位）列')
    design=spec.get('design')
    require(design in [None,'independent','paired','repeated_measures'],'design 只能是 independent/paired/repeated_measures')
    if mode=='raw_paired': require(design in [None,'paired'],'raw_paired 的 design 只能是 paired')
    elif mode in TIME: require(design in [None,'repeated_measures'],'时间曲线的 design 只能是 repeated_measures')
    else: require(design!='paired','配对数据请用 raw_paired')
    require(spec.get('sheet') is None or isinstance(spec.get('sheet'),str),'sheet 必须是工作表名')
    finite(spec.get('width_mm',165),'width_mm',80,400)
    finite(spec.get('height_mm',115),'height_mm',60,400)
    require(isinstance(spec.get('colors',[]),list),'colors 必须是列表')
    for c in spec.get('colors',[]): color(c)
    missing=spec.get('missing_values',DEFAULT_MISSING)
    require(isinstance(missing,list) and all(isinstance(v,str) for v in missing),'missing_values 是字符串列表')
    questions=[]
    if mode in UNIT_MODES:
        if not ('sample_id' in cols or layout=='wide') or not spec.get('replication_unit') or spec.get('independence_confirmed') is not True:
            if mode=='raw_paired':
                questions.append('配对单位是什么（受试者/动物/供体）？哪一列是配对 ID（宽表也可确认“同一行=同一单位”）？不同单位之间是否相互独立？')
            elif mode in TIME:
                questions.append('每条曲线代表的受试单位是什么（小鼠/患者/培养孔）？哪一列是其 ID？不同单位之间是否相互独立？')
            else:
                questions.append('独立实验单位是什么？每行是否为独立样本，并可提供 sample_id 列？技术重复请先确认合并规则。')
    if set(cols)&{'error','lower','upper'}:
        if not spec.get('error_definition'):
            questions.append('误差列/上下界表示 SD、SEM、CI 或什么范围？请提供定义与计算尺度。')
        if not cols.get('n') or not spec.get('replication_unit'):
            questions.append('各组样本量 n 和重复单位是什么？若无法确认，请去掉误差映射，只显示已有点估计并标注限制。')
        if spec.get('transform','none')!='none':
            raise CliError('已有误差的数据不自动 log2；请提供同尺度的变换后估计与区间')
    require(not(raw and spec.get('error_definition')),'raw 模式误差由 mode 明确；不接收自定义 error_definition')
    return questions[:3]


def records(spec, rows):
    """Long or wide input -> one record per observed cell, keeping row/column provenance."""
    cols=spec['columns'];out=[];blank=0
    if spec.get('layout','long')=='wide':
        for rowno,row in enumerate(rows,2):
            for c in cols['groups']:
                if row[c].strip()=='':
                    blank+=1;continue
                if 'sample_id' in cols: sid=row[cols['sample_id']]
                elif spec['mode']=='raw_paired': sid=f'row {rowno}'
                else: sid=f'{c} · row {rowno}'
                out.append({'row':rowno,'column':c,'group':c,'value':row[c],'sample_id':sid})
        return out,blank
    for rowno,row in enumerate(rows,2):
        if all(not str(v).strip() for v in row.values()):
            blank+=1;continue
        out.append({'row':rowno,'column':cols['value'],**{k:row[v] for k,v in cols.items()}})
    return out,blank


def tukey(vals):
    q1,med,q3=statistics.quantiles(vals,n=4,method='inclusive')
    iqr=q3-q1
    inside=[v for v in vals if q1-1.5*iqr<=v<=q3+1.5*iqr]
    return {'median':med,'q1':q1,'q3':q3,'iqr':iqr,'whisker_low':min(inside),'whisker_high':max(inside),
            'beyond_whiskers':len(vals)-len(inside)}


def spread(vals):
    sd=statistics.stdev(vals)
    return {'n':len(vals),'mean':statistics.mean(vals),'sd':sd,'sem':sd/math.sqrt(len(vals))}


def prepare_full(spec, rows):
    cols=spec['columns'];mode=spec['mode'];raw=mode.startswith('raw_')
    missing=spec.get('missing_values',DEFAULT_MISSING);tech=spec.get('technical_replicates','none')
    recs,blank=records(spec,rows);require(bool(recs),'没有可用数据行')
    clean=[];groups=[]
    for r in recs:
        rowno=r['row']
        g=text_value(r['group'].strip(),f'第 {rowno} 行 group',100)
        if g not in groups: groups.append(g)
        v=number(r['value'],f'第 {rowno} 行 {r["column"]}',missing)
        value=v
        if v is not None and spec.get('transform','none')=='log2':
            require(v>0,'log2 需要所有非缺失值 > 0；不自动增加伪计数')
            value=math.log2(v)
        item={'source_row':rowno,'source_column':r['column'],'group':g,'input_value':v,'plot_value':value}
        if 'sample_id' in r: item['sample_id']=text_value(str(r['sample_id']).strip(),f'第 {rowno} 行 sample_id',180)
        if 'x' in r: item['x']=number(r['x'],f'第 {rowno} 行 x',missing,False)
        for key in ['error','lower','upper','n']:
            if key in r:
                z=number(r[key],key,missing,False)
                if key=='n': require(z>=1 and z.is_integer(),'n 必须是正整数')
                if key=='error': require(z>=0,'误差不能是负数')
                item[key]=int(z) if key=='n' else z
        if 'lower' in item:
            require(v is not None and item['lower']<=v<=item['upper'],'区间需要 lower <= value <= upper')
        clean.append(item)
    order=spec.get('group_order',groups)
    require(isinstance(order,list) and len(order)==len(set(order)) and set(order)==set(groups),'group_order 必须恰好包含所有组')
    require(1<=len(order)<=MAX_GROUPS,f'每图支持 1—{MAX_GROUPS} 组；更多组请分面或明确额外设计')
    if not raw:
        if 'x' in cols:
            dup=[k for k,c in Counter((r['group'],r['x']) for r in clean).items() if c>1]
            require(not dup,'同组 x 重复；已有汇总不自动聚合时间点。逐个体原始数据请用 raw_timecourse_sd/sem')
        summary=[]
        for g in order:
            gr=[r for r in clean if r['group']==g]
            require(any(r['plot_value'] is not None for r in gr),f'{g} 全部缺失，不自动隐藏组')
            if mode!='line':
                require(len(gr)==1,'已有汇总每组只能一行；不自动平均已有结果')
                summary.append(gr[0])
        return {'clean':clean,'units':[],'summary':summary,'order':order,'blank_cells':blank,'shared_ids':[]}

    key=(lambda r:(r['group'],r['sample_id'],r.get('x'))) if 'sample_id' in clean[0] else (lambda r:(r['source_row'],r['source_column']))
    counts=Counter(key(r) for r in clean)
    dup=[k for k,c in counts.items() if c>1]
    if dup and tech!='mean':
        where='同一时间点' if mode in TIME else ''
        raise CliError(f'同组 sample_id 重复（{dup[0][0]} / {dup[0][1]}{where}）：若为技术重复（复孔、重复测读），'
                       '设置 technical_replicates: mean 先按生物学单位取均值；否则请确认数据结构。不把重复行当独立样本计数')
    units={}
    for r in clean:
        u=units.setdefault(key(r),{'group':r['group'],'sample_id':r.get('sample_id',f'row {r["source_row"]}'),
                                   **({'x':r['x']} if 'x' in r else {}),'values':[],'source_cells':[]})
        u['values'].append(r['plot_value']);u['source_cells'].append(f'{r["source_row"]}:{r["source_column"]}')
    unit_rows=[]
    for u in units.values():
        present=[v for v in u.pop('values') if v is not None]
        unit_rows.append({**u,'plot_value':statistics.mean(present) if present else None,
                          'technical_n':len(present),'technical_missing':len(u['source_cells'])-len(present),
                          'source_cells':' '.join(u['source_cells'])})
    ids={}
    for u in unit_rows: ids.setdefault(u['sample_id'],set()).add(u['group'])
    shared=sorted(s for s,gs in ids.items() if len(gs)>1 and not s.startswith('row ') and ' · row ' not in s)
    if mode in TIME:
        require(not shared,f'同一受试单位出现在多个组（如 {shared[:1]}）：交叉/组内设计本版不画为分组时间曲线')

    summary=[]
    for g in order:
        gu=[u for u in unit_rows if u['group']==g]
        vals=[u['plot_value'] for u in gu if u['plot_value'] is not None]
        require(bool(vals),f'{g} 全部缺失，不自动隐藏组')
        if mode in TIME:
            for x in sorted({u['x'] for u in gu}):
                xv=[u['plot_value'] for u in gu if u['x']==x and u['plot_value'] is not None]
                row={'group':g,'x':x,'n':len(xv),'mean':statistics.mean(xv) if xv else None}
                if len(xv)>=2: row.update(spread(xv))
                summary.append(row)
            continue
        agg={'group':g,'n':len(vals),'missing_units':len(gu)-len(vals),
             'observed_rows':sum(1 for r in clean if r['group']==g and r['plot_value'] is not None),
             'missing_rows':sum(1 for r in clean if r['group']==g and r['plot_value'] is None)}
        if mode in {'raw_mean_sd','raw_mean_sem','raw_paired'}:
            require(len(vals)>=2,'SD/SEM 至少需要两个已确认独立观测')
            agg.update(spread(vals))
        if mode=='raw_box':
            require(len(vals)>=5,f'箱线图每组至少需要 5 个独立单位（{g} 只有 {len(vals)}）；请改用 raw_points 或 raw_mean_sd')
            agg.update(tukey(vals))
        summary.append(agg)
    if mode in TIME:
        require(any(s['n']>=2 for s in summary),'时间曲线的 SD/SEM 至少需要某时间点有两个受试单位')
    if mode=='raw_paired':
        require(len(order)>=2,'配对图至少需要两个条件')
        value={(u['sample_id'],u['group']):u['plot_value'] for u in unit_rows}
        for s in summary[1:]:
            d=[value[(sid,s['group'])]-value[(sid,order[0])] for sid in ids
               if value.get((sid,s['group'])) is not None and value.get((sid,order[0])) is not None]
            require(len(d)>=2,f'{s["group"]} 与 {order[0]} 的完整配对少于 2 个')
            s.update({'paired_vs':order[0],'n_pairs':len(d),'mean_difference':statistics.mean(d),
                      'sd_difference':statistics.stdev(d)})
    return {'clean':clean,'units':unit_rows,'summary':summary,'order':order,'blank_cells':blank,'shared_ids':shared}


def prepare(spec, rows):
    p=prepare_full(spec,rows)
    return p['clean'],p['summary'],p['order']


def csv_out(path, records):
    keys=list(dict.fromkeys(k for r in records for k in r))
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=keys);writer.writeheader();writer.writerows(records)


def style_axes(ax):
    """Restrained journal axes after SciencePlots' rules: thin dark spines, outward ticks, no grid."""
    ax.spines[['top','right']].set_visible(False)
    ax.spines[['bottom','left']].set_color(INK)
    ax.spines[['bottom','left']].set_linewidth(.8)
    ax.tick_params(direction='out',length=3.5,width=.8,color=INK,labelcolor=INK)


def jitter(n, width=.13):
    import numpy as np
    return np.linspace(-width,width,n) if n>1 else np.array([0.])


def footer_text(spec, mode, summary, unit_name, technical):
    en=spec['language']=='en'
    if mode in TIME:
        stat='SD' if mode=='raw_timecourse_sd' else 'SEM'
        ns=[s['n'] for s in summary];span=f'{min(ns)}' if min(ns)==max(ns) else f'{min(ns)}–{max(ns)}'
        t=(f'Mean ± {stat} across units (unit: {unit_name}); thin lines: individual units; n per time point = {span}' if en
           else f'均值 ± {stat}（跨{unit_name}）；细线为单个{unit_name}；每时间点 n = {span}')
        if any(n < 2 for n in ns):
            t+=('; n=0: gap; n=1: mean only' if en else '；n=0：断点；n=1：仅均值')
    elif mode=='raw_paired':
        ns=', '.join(str(s['n_pairs']) for s in summary[1:])
        t=(f'Lines join the same {unit_name}; bar: mean; complete pairs vs {summary[0]["group"]} = {ns}' if en
           else f'连线为同一{unit_name}；横线为均值；与 {summary[0]["group"]} 的完整配对数 = {ns}')
    elif mode=='raw_box':
        ns=', '.join(str(s['n']) for s in summary)
        t=(f'Box: median and IQR; whiskers: most extreme values within 1.5 × IQR; all points shown; n = {ns}' if en
           else f'箱：中位数与四分位距；须：1.5×IQR 内最远值；显示全部点；各组 n = {ns}')
    elif mode in {'raw_mean_sd','raw_mean_sem'}:
        ns=', '.join(str(s['n']) for s in summary);stat='SD' if mode=='raw_mean_sd' else 'SEM'
        t=(f'Mean ± {stat}; n = {ns} (group order)' if en else f'均值 ± {stat}；各组 n = {ns}（按横轴顺序）')
    elif mode=='raw_points':
        t='Observed rows; no error bars' if en else '逐行观测；未绘制误差条'
    else:
        t=spec.get('error_definition','Supplied estimates; no uncertainty shown' if en else '用户已有估计；未提供误差')
    if technical:
        t+=(f'; technical replicates averaged within each {unit_name}' if en else f'；技术重复已在每个{unit_name}内取均值')
    return t


def render(spec, data_path, folder, changes=None, parent=None):
    header,rows=load_table(data_path,spec.get('sheet'))
    questions=validate(spec,header)
    original=checked_input_file(data_path,max_bytes=20_000_000)
    input_name='input'+original.suffix.lower()
    if questions:
        shutil.copyfile(original,folder/input_name)
        return pending(folder,spec,questions)
    p=prepare_full(spec,rows)
    clean,units,summary,order=p['clean'],p['units'],p['summary'],p['order']
    mode=spec['mode'];tech=spec.get('technical_replicates','none')=='mean'
    if p['shared_ids'] and mode in {'raw_mean_sd','raw_mean_sem','raw_box'} and spec.get('design')!='independent':
        shutil.copyfile(original,folder/input_name)
        return pending(folder,spec,[f'同一 sample_id 出现在多个组（如 {"、".join(p["shared_ids"][:3])}）：这是同一受试单位的配对/重复测量吗？'
                                    '若是请改用 raw_paired；若只是编号重名且各组相互独立，请设置 design: independent。'])
    plt=matplotlib_setup()
    import numpy as np
    from matplotlib import ticker
    fig,ax=plt.subplots(figsize=(spec.get('width_mm',165)/25.4,spec.get('height_mm',115)/25.4),layout='constrained')
    colors=spec.get('colors') or PALETTE
    require(len(colors)>=len(order),'colors 数量小于组数')
    style_axes(ax)
    unit_name=spec.get('replication_unit') or ('sample' if spec['language']=='en' else '样本')
    if mode in TIME:
        for i,g in enumerate(order):
            for sid in dict.fromkeys(u['sample_id'] for u in units if u['group']==g):
                tr=sorted([u for u in units if u['sample_id']==sid and u['group']==g],key=lambda u:u['x'])
                ax.plot([u['x'] for u in tr],[np.nan if u['plot_value'] is None else u['plot_value'] for u in tr],
                        color=colors[i],lw=.7,alpha=.35,zorder=2,gid=f'subject-{i}')
            gs=[s for s in summary if s['group']==g]
            stat='sd' if mode=='raw_timecourse_sd' else 'sem'
            ax.plot([s['x'] for s in gs],[np.nan if s['mean'] is None else s['mean'] for s in gs],color=colors[i],
                    marker=MARKS[i],ms=5,lw=1.6,ls=LINES[i],label=g,zorder=4,gid=f'mean-{i}')
            errors=[s for s in gs if stat in s]
            if errors:
                ax.errorbar([s['x'] for s in errors],[s['mean'] for s in errors],yerr=[s[stat] for s in errors],
                            fmt='none',color=colors[i],capsize=3,elinewidth=1,zorder=4)
        ax.legend(frameon=False,fontsize=9)
    elif mode.startswith('raw_'):
        pos={g:i for i,g in enumerate(order)}
        if mode=='raw_box':
            stats=[{'med':s['median'],'q1':s['q1'],'q3':s['q3'],'whislo':s['whisker_low'],'whishi':s['whisker_high'],
                    'fliers':[],'label':s['group']} for s in summary]
            b=ax.bxp(stats,positions=range(len(order)),widths=.52,showfliers=False,patch_artist=True,
                     medianprops={'color':INK,'lw':1.6},whiskerprops={'color':INK,'lw':.9},
                     capprops={'color':INK,'lw':.9},boxprops={'edgecolor':INK,'lw':.9})
            for i,patch in enumerate(b['boxes']): patch.set_facecolor(colors[i]);patch.set_alpha(.22)
        if mode=='raw_paired':
            for sid in dict.fromkeys(u['sample_id'] for u in units):
                tr=sorted([u for u in units if u['sample_id']==sid and u['plot_value'] is not None],key=lambda u:pos[u['group']])
                if len(tr)>1: ax.plot([pos[u['group']] for u in tr],[u['plot_value'] for u in tr],color='#9AA5B1',lw=.8,zorder=2)
        for i,g in enumerate(order):
            gu=[u for u in units if u['group']==g and u['plot_value'] is not None]
            offs=np.zeros(len(gu)) if mode=='raw_paired' else jitter(len(gu))
            if tech:
                for o,u in zip(offs,gu):
                    reps=[r['plot_value'] for r in clean if r['group']==g and r.get('sample_id')==u['sample_id']
                          and r['plot_value'] is not None]
                    ax.scatter([i+o]*len(reps),reps,s=9,facecolors='none',edgecolors=colors[i],linewidth=.6,alpha=.6,zorder=2)
            ax.scatter(i+offs,[u['plot_value'] for u in gu],c=colors[i],marker=MARKS[i],s=33,edgecolors='white',
                       linewidth=.5,zorder=3,gid='observations-'+str(i))
            s=summary[i]
            if mode in {'raw_mean_sd','raw_mean_sem'}:
                ax.errorbar(i,s['mean'],yerr=s['sd'] if mode=='raw_mean_sd' else s['sem'],fmt='_',ms=21,lw=1.4,
                            c=INK,capsize=5,zorder=4)
            if mode=='raw_paired':
                ax.plot([i-.18,i+.18],[s['mean']]*2,color=INK,lw=2.2,solid_capstyle='butt',zorder=4)
        ax.set_xticks(range(len(order)),order)
        ax.set_xlim(-.5,len(order)-.5)
    elif mode in {'summary_bar','summary_point'}:
        for i,s in enumerate(summary):
            y=s['plot_value']
            require(y is not None,'汇总估计缺失；请明确缺失展示设计')
            err=None
            if 'error' in s: err=s['error']
            if 'lower' in s: err=[[y-s['lower']],[s['upper']-y]]
            if mode=='summary_bar': ax.bar(i,y,color=colors[i],width=.55,edgecolor=INK,lw=.6,zorder=2,hatch=HATCH[i])
            else: ax.scatter(i,y,c=colors[i],marker=MARKS[i],s=48,zorder=3)
            if err is not None: ax.errorbar(i,y,yerr=err,fmt='none',c=INK,capsize=4,lw=1.2,zorder=4)
        ax.set_xticks(range(len(order)),order);ax.set_xlim(-.5,len(order)-.5)
        if mode=='summary_bar':
            lo,hi=ax.get_ylim();ax.set_ylim(min(0,lo),max(0,hi))
    else:
        for i,g in enumerate(order):
            gr=sorted([r for r in clean if r['group']==g],key=lambda r:r['x'])
            xs=[r['x'] for r in gr];ys=[np.nan if r['plot_value'] is None else r['plot_value'] for r in gr]
            ax.plot(xs,ys,color=colors[i],marker=MARKS[i],ls=LINES[i],label=g,lw=1.4,ms=5)
            if 'error' in spec['columns']:
                ax.errorbar(xs,ys,yerr=[r['error'] for r in gr],fmt='none',c=colors[i],capsize=3)
            if 'lower' in spec['columns']:
                ax.errorbar(xs,ys,yerr=[[y-r['lower'] for y,r in zip(ys,gr)],
                                       [r['upper']-y for y,r in zip(ys,gr)]],fmt='none',c=colors[i],capsize=3)
        ax.legend(frameon=False,fontsize=9)
    unit=spec['unit'];transform=spec.get('transform','none')
    default_y=('log2'+' ('+unit+')') if transform=='log2' else unit
    default_x=spec['columns'].get('x') if mode in TIME else ('分组' if spec['language']=='zh' else 'Group')
    ax.set(xlabel=spec.get('xlabel',default_x),ylabel=spec.get('ylabel',default_y),title=spec['title'])
    ax.title.set_fontsize(13)
    ax.yaxis.set_major_locator(ticker.MaxNLocator(nbins=6))
    error_label={'raw_points':'不计算误差；逐行观测','raw_mean_sd':'均值 ± 样本 SD (ddof=1)',
                 'raw_mean_sem':'均值 ± SEM = SD/√n','raw_box':'Tukey 箱线图：中位数、四分位数（线性插值）、1.5×IQR 须；不计算误差条',
                 'raw_paired':'逐单位配对连线 + 各条件均值；不绘误差条；差值仅描述性汇总',
                 'raw_timecourse_sd':'每时间点跨受试单位的均值 ± 样本 SD (ddof=1)',
                 'raw_timecourse_sem':'每时间点跨受试单位的均值 ± SEM = SD/√n（n 为该时间点单位数）'}
    uncertainty=error_label.get(mode,spec.get('error_definition','未显示误差；仅已有估计'))
    fig.supxlabel(footer_text(spec,mode,summary,unit_name,tech),fontsize=8,color='#526477')
    raw=mode.startswith('raw_')
    caption={'purpose':spec['purpose'],'unit':unit,'group_order':order,'mode':mode,
             'uncertainty':uncertainty,'replication_unit':spec.get('replication_unit','未确认'),
             'transform': 'log2(value)，无伪计数；误差在变换后尺度计算' if transform=='log2' else 'none',
             'missing_rows':sorted({r['source_row'] for r in clean if r['plot_value'] is None}),
             'missing_policy':'保留原表和记录；原始点只绘制非缺失值；line 和时间曲线保留明确缺失断点；时间曲线 n=0 保留空均值，n=1 无 SD/SEM；不存在的时间点无法识别',
             'statistical_tests':'未执行；未生成 p 值或显著性标记','notes':spec.get('notes',[])}
    if raw:
        caption.update({'layout':spec.get('layout','long'),'design':spec.get('design') or
                        {'raw_paired':'paired','raw_timecourse_sd':'repeated_measures','raw_timecourse_sem':'repeated_measures'}.get(mode,'independent'),
                        'center':{'raw_box':'median','raw_points':'none'}.get(mode,'mean'),
                        'spread':{'raw_mean_sd':'SD','raw_mean_sem':'SEM','raw_timecourse_sd':'SD','raw_timecourse_sem':'SEM',
                                  'raw_box':'IQR + 1.5×IQR whiskers'}.get(mode,'none'),
                        'n_definition':'独立单位数（sample_id）；技术重复不计入 n' if 'sample_id' in spec['columns'] or spec.get('layout')=='wide' else '行数',
                        'n_per_group':{s['group']:s['n'] for s in summary} if mode not in TIME else
                        {g:{str(s['x']):s['n'] for s in summary if s['group']==g} for g in order}})
        if tech:
            tn=[u['technical_n'] for u in units]
            caption['technical_replicates']={'rule':'先逐值变换，再在每个 sample_id 内取均值；统计只用单位均值',
                                             'per_unit_range':[min(tn),max(tn)],'excluded_missing':sum(u['technical_missing'] for u in units)}
        if spec.get('layout')=='wide':
            caption['wide_mapping']={'group_columns':spec['columns']['groups'],'blank_cells_skipped':p['blank_cells'],
                                     'unit_key':spec['columns'].get('sample_id') or ('source row' if mode=='raw_paired' else 'column · row（各格视为不同单位）')}
        if mode=='raw_paired':
            caption['paired_differences']=[{k:s[k] for k in ['group','paired_vs','n_pairs','mean_difference','sd_difference']}
                                           for s in summary[1:]]
            caption['paired_note']=('差值 = 条件 − 首个条件，只用完整配对，仅描述性；未执行配对检验。'
                                    '各条件均值横线使用该条件全部已观测单位')
    qa=text_checks(fig)
    export(fig,folder,{'kind':'results','input':input_name,'input_sha256':digest(original),
                       'source':'source.json','caption':'caption.json','transform':caption['transform'],
                       'uncertainty':uncertainty,'group_order':order})
    plt.close(fig)
    shutil.copyfile(original,folder/input_name)
    csv_out(folder/'plot-data.csv',clean);csv_out(folder/'summary.csv',summary or clean)
    if raw: csv_out(folder/'units.csv',units)
    write_json(folder/'source.json',spec);write_json(folder/'caption.json',caption)
    if changes is not None: write_json(folder/'changes.json',{'parent':parent,'changes':changes})
    from common import package_replay
    package_replay(folder,'results',input_name)
    manifest={'status':'GENERATED_REVIEW_PENDING','kind':'results','parent':parent,
              'input_sha256':digest(original),'editability':{
                  'text':'figure.svg independent text; requires bundled font',
                  'data_and_styles':'source.json + input table + reproduce.py',
                  'geometry':'SVG vector paths; changes to values require rerender',
                  'raster_assets':[]}}
    notes=('已生成真实数据图、SVG/PNG/PDF、原始输入副本、plot-data.csv（逐格来源）、'+('units.csv（统计单位）、' if raw else '')+
           'summary.csv 和 caption.json。\n误差：'+uncertainty+'；变换：'+caption['transform']+'。\n'
           '未执行统计检验或生信分析，未生成 p 值。变量、单位、分组、缺失、n 与重复单位详见 caption.json。\n'
           '运行 reproduce.py 可在 rerun 新版本复现；所需脚本、字体和依赖说明随附。\n'
           'SVG 文字独立、几何为矢量；颜色与标签可局部修改；outlined 版本文字转轮廓。科学复核待完成。')
    if qa['issues']: notes+='\n边界筛查发现问题，详见 qa.json；使用前先修正布局。'
    return finish(folder,manifest,notes,qa)


ID_RE=re.compile(r'(?i)(\bid\b|_id$|sample|subject|patient|mouse|mice|animal|donor|\brat\b|replicate|\brep\b|biolog|'
                 r'样本|编号|受试|患者|病人|小鼠|大鼠|动物|供体|个体|生物学重复)')
TECH_RE=re.compile(r'(?i)(technical|\bwell\b|\btech|复孔|技术重复|孔位)')
TIME_RE=re.compile(r'(?i)(time|\bday|\bhour|\bweek|\bmin|\bh\b|时间|天|周|小时|分钟|时点)')
GROUP_RE=re.compile(r'(?i)(group|treatment|condition|\barm\b|genotype|cohort|dose|组|处理|条件|基因型|剂量)')


def profile(header, rows, missing=None):
    """Read-only structure guess for inspect-data; the host confirms before rendering."""
    missing=missing or DEFAULT_MISSING
    info=[]
    for c in header:
        vals=[r[c].strip() for r in rows]
        present=[v for v in vals if v not in missing]
        nums=[]
        for v in present:
            try: nums.append(float(v))
            except ValueError: pass
        info.append({'name':c,'filled':len(present),'missing':sum(1 for v in vals if v and v in missing),
                     'blank':sum(1 for v in vals if not v),'numeric':bool(present) and len(nums)==len(present),
                     'distinct':len(set(present)),'examples':list(dict.fromkeys(present))[:3]})
    by={i['name']:i for i in info}
    unit=None
    numeric=[i['name'] for i in info if i['numeric']]
    text=[i['name'] for i in info if not i['numeric'] and i['filled']]
    n=len(rows)
    idc=next((c for c in text+numeric if ID_RE.search(c) and not TECH_RE.search(c) and not GROUP_RE.search(c)),None)
    techc=next((c for c in header if TECH_RE.search(c)),None)
    timec=next((c for c in numeric if TIME_RE.search(c) and by[c]['distinct']<=min(30,n/2) and c!=idc),None)
    groupc=next((c for c in text if c not in (idc,techc) and 2<=by[c]['distinct']<=MAX_GROUPS and GROUP_RE.search(c)),None) \
        or next((c for c in text if c not in (idc,techc) and 2<=by[c]['distinct']<=MAX_GROUPS and by[c]['distinct']<n),None)
    values=[c for c in numeric if c not in (idc,timec,techc)]
    out={'status':'PROFILED','rows':n,'columns':info,'statistical_tests':'不执行；只识别结构'}
    questions=[];spec={'kind':'results','language':'zh','title':'待确认标题','purpose':'待确认用途'}
    if groupc and values:
        out['layout_guess']='long'
        value=values[0] if len(values)==1 else None
        if value is None: questions.append('有多个数值列（'+'、'.join(values[:5])+'）：要画哪一列？若它们是各组测量值，这是宽表。')
        cols={'group':groupc,'value':value or values[0]}
        if idc: cols['sample_id']=idc
        if timec: cols['x']=timec
        m=re.search(r'[(（]([^()（）]+)[)）]\s*$',cols['value']);unit=m.group(1) if m else None
        per=Counter((r[groupc],r[idc] if idc else i,r[timec] if timec else None) for i,r in enumerate(rows))
        tech=idc is not None and max(per.values())>1
        groups_of={}
        if idc:
            for r in rows: groups_of.setdefault(r[idc],set()).add(r[groupc])
        shared=[s for s,g in groups_of.items() if len(g)>1]
        out['structure']={'group_column':groupc,'value_column':cols['value'],'id_column':idc,'time_column':timec,
                          'technical_column':techc,'groups':list(dict.fromkeys(r[groupc] for r in rows)),
                          'rows_per_unit_max':max(per.values()),'ids_shared_across_groups':len(shared)}
        if timec and idc and not shared:
            spec['mode']='raw_timecourse_sem';spec['design']='repeated_measures'
            questions.append(f'每条曲线是一个“{idc}”吗？误差用 SD（个体变异）还是 SEM（均值精度）？默认 SEM，不做组间检验。')
        elif timec and not idc:
            spec['mode']='line'
            if max(per.values())>1: questions.append('同一组同一时间点有多行，但没有受试单位 ID 列：哪一列标识个体？')
        elif shared:
            spec['mode']='raw_paired';spec['design']='paired'
            questions.append(f'同一“{idc}”出现在多个组：是同一受试单位的配对/重复测量吗？若是将画配对连线，不做检验。')
        else:
            ns=Counter(r[groupc] for r in rows if r[cols['value']].strip() not in missing)
            spec['mode']='raw_box' if min(ns.values())>=8 and not tech else 'raw_mean_sd'
            questions.append(f'各组的独立实验单位是什么（{idc or "每行"}）？各单位之间是否相互独立？')
        if tech:
            spec['technical_replicates']='mean'
            questions.insert(0,f'同一“{idc}”有多行{"（"+techc+"）" if techc else ""}：是技术重复吗？若是先按单位取均值，n 按单位计。')
    elif len(values)>=2:
        out['layout_guess']='wide'
        cols={'groups':values}
        if idc: cols['sample_id']=idc
        paired_hint=idc is not None
        out['structure']={'group_columns':values,'id_column':idc,'blank_cells':sum(by[c]['blank'] for c in values)}
        spec['layout']='wide'
        if paired_hint:
            spec['mode']='raw_paired';spec['design']='paired'
            questions.append(f'每行是同一“{idc}”在各列条件下的测量吗（配对）？若各列是不同个体，请说明。')
        else:
            ns=[by[c]['filled'] for c in values]
            spec['mode']='raw_box' if min(ns)>=8 else 'raw_mean_sd'
            questions.append('每列是一个组、每个格子是一个独立单位吗？同一行是否为同一受试单位（若是则为配对）？')
        heads=[re.search(r'[(（]([^()（）]+)[)）]\s*$',c) for c in values]
        unit=heads[0].group(1) if heads[0] and all(h and h.group(1)==heads[0].group(1) for h in heads) else None
    else:
        out['layout_guess']='unknown'
        cols={}
        questions.append('没有找到“分组列 + 数值列”或“多列数值”的结构：哪一列是分组，哪一列是测量值？')
    spec['columns']=cols
    if unit: spec['unit']=unit
    else: questions.append('测量值的单位是什么？')
    spec['replication_unit']='待确认';spec['independence_confirmed']=False
    out['suggested_spec']=spec
    out['questions']=questions[:3]
    out['next']='确认问题后补全 title/purpose/unit/replication_unit/independence_confirmed，再运行 b2r.py results'
    return out
